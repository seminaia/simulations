#!/usr/bin/env python3
"""
ellingham_v6_data.py

Turns the raw, wide `TP-data` sheet of Data_EllinghamDiagram_v6.xlsx into a
clean, readable workbook:  data/ellingham/ellingham_v6_clean.xlsx

Raw layout (hard to read): ~270 columns, a pair 'x#NAME' / 'y#NAME' per reaction
    x = 1000 / T[K]
    y = equilibrium pO2 in torr (0 where the equilibrium is not defined)

Clean layout:
  README     what every sheet / column means
  Reactions  one row per reaction: metal, balanced reaction, T range, dG range
  Curves     long ("tidy") table: one row per (reaction, temperature)

Free energy convention (same as the Ellingham diagram):
    dG = R * T * ln(pO2 / 1 atm)    in kJ per mole of O2, pO2 [atm] = torr / 760
so every reaction is written with exactly one O2 on the left-hand side.

Usage:
    python ellingham_v6_data.py          # (re)build the clean workbook
"""
import re
from fractions import Fraction
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
RAW_WORKBOOK = HERE / 'data' / 'ellingham' / 'Data_EllinghamDiagram_v6.xlsx'
RAW_SHEET = 'TP-data'
CLEAN_WORKBOOK = HERE / 'data' / 'ellingham' / 'ellingham_v6_clean.xlsx'
R_KJ = 8.314462618e-3  # kJ/(mol K)
TORR_PER_ATM = 760.0   # raw pO2 is in torr; reference state is 1 atm (101325 Pa)

NUMBER_FORMATS = {
    'T_K': '0.0', 'T_C': '0.0', 'pO2_torr': '0.000E+00', 'pO2_atm': '0.000E+00',
    'log10_pO2_atm': '0.000',
    'dG_kJ_per_molO2': '0.00', 'T_min_K': '0.0', 'T_max_K': '0.0',
    'dG_at_Tmin_kJ': '0.00', 'dG_at_Tmax_kJ': '0.00',
}


# ----------------------------------------------------------------------
# Reaction names -> balanced reaction strings
# ----------------------------------------------------------------------
def _parse(part):
    """'FE3O4' / 'O12TB7' / 'MO1O2.875' -> {'Fe': 3, 'O': 4} (Fractions)."""
    return {el.capitalize(): Fraction(n)
            for el, n in re.findall(r'([A-Z]+)(\d+(?:\.\d+)?)', part)}


def _num(n):
    return '' if n == 1 else (str(int(n)) if n.denominator == 1 else f'{float(n):g}')


def _formula(comp):
    metals = ''.join(f'{e}{_num(n)}' for e, n in comp.items() if e != 'O')
    return f"{metals}O{_num(comp['O'])}"


def _coef(x):
    x = Fraction(x).limit_denominator(10000)
    if x == 1:
        return ''
    return str(x.numerator) if x.denominator == 1 else f'({x.numerator}/{x.denominator})'


def describe(name):
    """Raw column name -> (metal, lower_oxide, higher_oxide, reaction) per mole O2.

    'FE1O1'       : 2Fe + O2 = 2FeO                (formation of one oxide)
    'FE3O4_FE2O3' : 4Fe3O4 + O2 = 6Fe2O3           (oxidation of a lower oxide)
    """
    comps = [_parse(p) for p in name.split('_')]
    metal = next(e for e in comps[0] if e != 'O')
    if len(comps) == 1:
        a = comps[0]
        m, n = a[metal], a['O']
        reaction = f'{_coef(2 * m / n)}{metal} + O2 = {_coef(2 / n)}{_formula(a)}'
        return metal, '', _formula(a), reaction
    a, b = comps[0], comps[1]
    a1, b1, a2, b2 = a[metal], a['O'], b[metal], b['O']
    den = a1 * b2 - b1 * a2
    x, y = 2 * a2 / den, 2 * a1 / den
    reaction = f'{_coef(x)}{_formula(a)} + O2 = {_coef(y)}{_formula(b)}'
    return metal, _formula(a), _formula(b), reaction


# ----------------------------------------------------------------------
# Raw workbook -> tidy tables
# ----------------------------------------------------------------------
def read_raw(path=RAW_WORKBOOK, sheet=RAW_SHEET):
    frame = pd.read_excel(path, sheet_name=sheet, header=0)

    pairs = {}
    for col in frame.columns:
        m = re.match(r'^([xy])(\d+)#(.+)$', str(col).strip())
        if m:
            pairs.setdefault(int(m.group(2)), {'name': m.group(3)})[m.group(1)] = col

    reactions, curves, seen = [], [], set()
    for idx in sorted(pairs):
        p = pairs[idx]
        name = p['name']
        if 'x' not in p or 'y' not in p or name in seen:
            continue
        x = pd.to_numeric(frame[p['x']], errors='coerce').to_numpy(float)
        po2 = pd.to_numeric(frame[p['y']], errors='coerce').to_numpy(float)
        ok = np.isfinite(x) & np.isfinite(po2) & (x > 0) & (po2 > 0)
        if ok.sum() < 2:
            continue
        seen.add(name)

        T = 1000.0 / x[ok]
        po2_torr = po2[ok]
        order = np.argsort(T)
        T, po2_torr = T[order], po2_torr[order]
        po2 = po2_torr / TORR_PER_ATM
        dG = R_KJ * T * np.log(po2)

        metal, lower, higher, reaction = describe(name)
        o_per_metal = _parse(name.split('_')[-1])
        reactions.append({
            'Reaction_ID': name, 'Metal': metal, 'Reaction': reaction,
            'Lower_oxide': lower, 'Higher_oxide': higher,
            'T_min_K': T[0], 'T_max_K': T[-1], 'N_points': len(T),
            'dG_at_Tmin_kJ': dG[0], 'dG_at_Tmax_kJ': dG[-1],
            '_ratio': float(o_per_metal['O'] / o_per_metal[metal]),
        })
        curves.append(pd.DataFrame({
            'Reaction_ID': name, 'Metal': metal, 'Reaction': reaction,
            'T_K': T, 'T_C': T - 273.15, 'pO2_torr': po2_torr, 'pO2_atm': po2, 'log10_pO2_atm': np.log10(po2),
            'dG_kJ_per_molO2': dG,
        }))

    reactions = (pd.DataFrame(reactions)
                 .sort_values(['Metal', '_ratio', 'Reaction_ID'])
                 .drop(columns='_ratio').reset_index(drop=True))
    order = {rid: i for i, rid in enumerate(reactions['Reaction_ID'])}
    curves = pd.concat(curves, ignore_index=True)
    curves = (curves.assign(_o=curves['Reaction_ID'].map(order))
              .sort_values(['_o', 'T_K']).drop(columns='_o').reset_index(drop=True))
    return reactions, curves


README = [
    ('Ellingham data — cleaned version of Data_EllinghamDiagram_v6.xlsx (sheet TP-data)', ''),
    ('', ''),
    ('Sheet', 'Contents'),
    ('Reactions', 'One row per reaction line on the diagram.'),
    ('Curves', 'One row per (reaction, temperature) point. Filter by Reaction_ID or Metal.'),
    ('', ''),
    ('Column', 'Meaning'),
    ('Reaction_ID', 'Original column name in the raw sheet (e.g. FE3O4_FE2O3).'),
    ('Metal', 'Element symbol of the metal.'),
    ('Reaction', 'Balanced reaction written per mole of O2. A single oxide is a formation '
                 'reaction (2Fe + O2 = 2FeO); two oxides is oxidation of the lower oxide '
                 '(4Fe3O4 + O2 = 6Fe2O3).'),
    ('Lower_oxide / Higher_oxide', 'Oxide on the left / right of the reaction (empty for formation from the metal).'),
    ('T_K, T_C', 'Temperature in kelvin and degrees Celsius (T_K = 1000 / x in the raw sheet).'),
    ('pO2_torr', 'Equilibrium oxygen partial pressure in torr, as in the raw column y. Points where it is 0 / undefined are dropped.'),
    ('pO2_atm', 'pO2_torr / 760.'),
    ('log10_pO2_atm', 'log10 of pO2_atm.'),
    ('dG_kJ_per_molO2', 'Standard free energy  dG = R * T * ln(pO2_atm)  in kJ per mole of O2 (R = 8.314462618 J/mol/K, reference 1 atm).'),
    ('T_min_K / T_max_K', 'Temperature range covered by the reaction.'),
    ('dG_at_Tmin_kJ / dG_at_Tmax_kJ', 'dG at the ends of that range.'),
]


def _style(ws):
    head_fill = PatternFill('solid', start_color='305496')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    headers = [c.value for c in ws[1]]
    for i, header in enumerate(headers, start=1):
        letter = get_column_letter(i)
        width = max(len(str(c.value)) for c in ws[letter][:200] if c.value is not None)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        fmt = NUMBER_FORMATS.get(header)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt


def build_clean_workbook(raw=RAW_WORKBOOK, out=CLEAN_WORKBOOK):
    print('-> Reading raw workbook...', flush=True)
    reactions, curves = read_raw(raw)
    print(f'-> Writing {out.name} (about {len(curves)} rows)...', flush=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine='openpyxl') as xw:
        pd.DataFrame(README).to_excel(xw, sheet_name='README', index=False, header=False)
        reactions.to_excel(xw, sheet_name='Reactions', index=False)
        curves.to_excel(xw, sheet_name='Curves', index=False)
        _style(xw.book['Reactions'])
        _style(xw.book['Curves'])
        readme = xw.book['README']
        readme.column_dimensions['A'].width = 32
        readme.column_dimensions['B'].width = 110
        for row in readme.iter_rows():
            row[0].font = Font(bold=True)
            row[1].alignment = Alignment(wrap_text=True, vertical='top')
    print(f'Wrote {out}  ({len(reactions)} reactions, {len(curves)} points)')
    return out


def ensure_clean_workbook(rebuild=False):
    """Build the clean workbook if it is missing or older than the raw file."""
    stale = (not CLEAN_WORKBOOK.exists()
             or CLEAN_WORKBOOK.stat().st_mtime < RAW_WORKBOOK.stat().st_mtime)
    if rebuild or stale:
        build_clean_workbook()
    return CLEAN_WORKBOOK


def load_clean(rebuild=False):
    """Return (reactions_df, curves_df) from the clean workbook."""
    path = ensure_clean_workbook(rebuild)
    return (pd.read_excel(path, sheet_name='Reactions').fillna(''),
            pd.read_excel(path, sheet_name='Curves'))


if __name__ == '__main__':
    build_clean_workbook()
